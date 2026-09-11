# -*- coding: utf-8 -*-
"""问题3：日前计划、日内调整与10分钟储能反馈控制。

实现口径：
1. 0:00生成全天初始计划g0；6/12/18点按策略更新未来交付区间；
2. 每个决策点用成熟的历史连续24小时残差轨迹构造SAA场景；
3. SAA中的储能动作是情景补救变量，真实动作由10分钟因果DP执行器决定；
4. 日内不重置SOC，全年从1月1日连续回放；
5. 每个时段的最终购电相对当天0:00的g0只结算一次。

本模型是“两阶段补救近似 + 因果滚动随机MPC”，不是严格多阶段全局最优。
"""

from __future__ import annotations

import argparse
import copy
import json
import math
import shutil
import time
from dataclasses import asdict, dataclass
from datetime import date, datetime, time as dt_time, timedelta
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.optimize import linprog
from scipy.sparse import coo_matrix


SLOTS_PER_DAY = 144
SLOT_MINUTES = 10
UPDATE_SLOTS = {6: 36, 12: 72, 18: 108}
YEAR_START = pd.Timestamp("2025-01-01 00:00:00")
YEAR_END = pd.Timestamp("2026-01-01 00:00:00")

STRATEGIES: dict[str, tuple[int, ...]] = {
    "S0": (),
    "S06": (6,),
    "S12": (12,),
    "S18": (18,),
    "S0612": (6, 12),
    "S0618": (6, 18),
    "S1218": (12, 18),
    "S061218": (6, 12, 18),
}
STRATEGY_ALIASES = {"S6": "S06", "Sall": "S061218"}
CUMULATIVE_ABLATION = ("S0", "S06", "S0612", "S061218")


@dataclass(frozen=True)
class StorageParams:
    soc_min: float = 1200.0
    soc_max: float = 10800.0
    soc_initial: float = 6000.0
    eta_charge: float = 0.9
    eta_discharge: float = 0.9
    max_power_kw: float = 5000.0
    delta_h: float = 1.0 / 6.0

    @property
    def max_charge_kwh(self) -> float:
        return self.max_power_kw * self.delta_h

    @property
    def max_discharge_kwh(self) -> float:
        return self.max_power_kw * self.delta_h


@dataclass(frozen=True)
class RunConfig:
    scenario_count: int = 12
    history_days: int = 30
    seed: int = 2025
    soc_step: float = 100.0
    emergency_multiplier: float = 5.0
    throughput_penalty: float = 1.0e-4
    feedback_terminal_penalty: float = 2.0
    final_soc: float = 6000.0
    solver_tolerance: float = 1.0e-7


@dataclass
class ProblemData:
    dates: pd.DatetimeIndex
    input_time_labels: list[str]
    price: np.ndarray
    reference_load_kwh: np.ndarray
    reference_pv_kwh: np.ndarray
    actual_load_kwh: np.ndarray
    actual_pv_kwh: np.ndarray
    forecasts: pd.DataFrame
    data_audit: dict

    def __post_init__(self) -> None:
        self._date_to_index = {
            pd.Timestamp(value).date(): idx for idx, value in enumerate(self.dates)
        }
        self._actual_load_flat = self.actual_load_kwh.reshape(-1)
        self._actual_pv_flat = self.actual_pv_kwh.reshape(-1)
        self._forecast_lookup = {
            pd.Timestamp(row.issue_time): row.hourly_kw.copy()
            for row in self.forecasts.itertuples(index=False)
        }

    def day_index(self, value: date) -> int:
        try:
            return self._date_to_index[value]
        except KeyError as exc:
            raise ValueError(f"附件2中不存在日期：{value}") from exc

    @staticmethod
    def slot_for_target(target_time: pd.Timestamp) -> int:
        minutes = target_time.hour * 60 + target_time.minute
        if minutes == 0:
            return SLOTS_PER_DAY - 1
        return minutes // SLOT_MINUTES - 1

    def actual_at(
        self,
        target_time: pd.Timestamp,
        kind: str,
        missing: float = np.nan,
    ) -> float:
        minutes = int((pd.Timestamp(target_time) - YEAR_START).total_seconds() // 60)
        if minutes <= 0 or minutes % SLOT_MINUTES != 0:
            return float(missing)
        index = minutes // SLOT_MINUTES - 1
        values = self._actual_load_flat if kind == "load" else self._actual_pv_flat
        if index < 0 or index >= len(values):
            return float(missing)
        return float(values[index])

    def actual_vector(self, targets: pd.DatetimeIndex, kind: str) -> np.ndarray:
        return np.asarray([self.actual_at(t, kind) for t in targets], dtype=float)


@dataclass
class ScenarioBundle:
    targets: pd.DatetimeIndex
    base_load: np.ndarray
    base_pv: np.ndarray
    load: np.ndarray
    pv: np.ndarray
    probabilities: np.ndarray
    source_issues: list[pd.Timestamp]
    decision_time: pd.Timestamp

    @property
    def net_load(self) -> np.ndarray:
        return self.load - self.pv


@dataclass
class PlanningResult:
    objective: float
    g0: np.ndarray | None
    committed_block: np.ndarray
    proposed_current_day: np.ndarray
    mean_soc: np.ndarray
    scenario_g: np.ndarray
    scenario_charge: np.ndarray
    scenario_discharge: np.ndarray
    solver_message: str


def _time_label(value) -> str:
    if isinstance(value, dt_time):
        return value.strftime("%H:%M:%S")
    return str(value)


def _numeric_matrix(frame: pd.DataFrame, name: str) -> np.ndarray:
    values = frame.apply(pd.to_numeric, errors="coerce").to_numpy(dtype=float)
    if np.isnan(values).any():
        raise ValueError(f"{name}存在缺失或非数值数据")
    if not np.isfinite(values).all():
        raise ValueError(f"{name}存在NaN或无穷值")
    if (values < 0).any():
        raise ValueError(f"{name}存在负数")
    return values


def read_problem_data(data_dir: Path) -> ProblemData:
    """读取附件1、2、3并执行结构审计；原始功率统一除以6得到kWh。"""

    attachment1 = data_dir / "附件1.xlsx"
    attachment2 = data_dir / "附件2.xlsx"
    attachment3 = data_dir / "附件3.xlsx"
    for path in (attachment1, attachment2, attachment3):
        if not path.exists():
            raise FileNotFoundError(f"缺少必需附件：{path}")

    ref = pd.read_excel(attachment1)
    if ref.shape != (SLOTS_PER_DAY, 4):
        raise ValueError(f"附件1应为144×4数据区，实际为{ref.shape}")
    ref.columns = [str(x).strip() for x in ref.columns]
    price_col = next((c for c in ref.columns if "电价" in c), None)
    load_col = next((c for c in ref.columns if "负载" in c or "负荷" in c), None)
    pv_col = next((c for c in ref.columns if "光伏" in c), None)
    if not all((price_col, load_col, pv_col)):
        raise ValueError(f"附件1字段无法识别：{list(ref.columns)}")
    price = _numeric_matrix(ref[[price_col]], "附件1电价").ravel()
    reference_load_kwh = _numeric_matrix(ref[[load_col]], "附件1负荷").ravel() / 6.0
    reference_pv_kwh = _numeric_matrix(ref[[pv_col]], "附件1光伏").ravel() / 6.0
    input_time_labels = [_time_label(x) for x in ref.iloc[:, 0].tolist()]

    sheets = pd.ExcelFile(attachment2).sheet_names
    expected_sheets = {"小区负载", "光伏发电实际功率"}
    if not expected_sheets.issubset(sheets):
        raise ValueError(f"附件2缺少工作表，实际为{sheets}")
    load_raw = pd.read_excel(attachment2, sheet_name="小区负载")
    pv_raw = pd.read_excel(attachment2, sheet_name="光伏发电实际功率")
    if load_raw.shape != (365, 145) or pv_raw.shape != (365, 145):
        raise ValueError(
            f"附件2两个数据区应为365×145，实际为{load_raw.shape}和{pv_raw.shape}"
        )
    dates = pd.DatetimeIndex(pd.to_datetime(load_raw.iloc[:, 0], errors="coerce"))
    pv_dates = pd.DatetimeIndex(pd.to_datetime(pv_raw.iloc[:, 0], errors="coerce"))
    if dates.isna().any() or pv_dates.isna().any() or not dates.equals(pv_dates):
        raise ValueError("附件2负荷和光伏日期缺失或不一致")
    expected_dates = pd.date_range("2025-01-01", "2025-12-31", freq="D")
    if not dates.equals(expected_dates):
        raise ValueError(
            f"附件2日期应为2025全年，实际为{dates.min()}至{dates.max()}"
        )
    actual_load_kwh = _numeric_matrix(load_raw.iloc[:, 1:], "附件2负荷") / 6.0
    actual_pv_kwh = _numeric_matrix(pv_raw.iloc[:, 1:], "附件2光伏") / 6.0

    forecast_raw = pd.read_excel(attachment3)
    if forecast_raw.shape != (1460, 26):
        raise ValueError(f"附件3应为1460×26数据区，实际为{forecast_raw.shape}")
    date_values = forecast_raw.iloc[:, 0].replace(r"^\s*$", np.nan, regex=True).ffill()
    forecast_dates = pd.to_datetime(date_values, errors="coerce")
    if forecast_dates.isna().any():
        raise ValueError("附件3日期向下填充后仍有无法识别值")
    issue_text = forecast_raw.iloc[:, 1].astype(str).str.strip()
    issue_hour = pd.to_numeric(issue_text.str.split(":").str[0], errors="coerce")
    if issue_hour.isna().any() or not set(issue_hour.astype(int).unique()).issubset({0, 6, 12, 18}):
        raise ValueError("附件3预报时刻必须是0:00、6:00、12:00、18:00")
    issue_times = forecast_dates + pd.to_timedelta(issue_hour.astype(int), unit="h")
    hourly_kw = _numeric_matrix(forecast_raw.iloc[:, 2:], "附件3光伏预测")
    forecasts = pd.DataFrame(
        {"issue_time": issue_times, "hourly_kw": [row for row in hourly_kw]}
    )
    if forecasts["issue_time"].duplicated().any():
        raise ValueError("附件3存在重复的预测发布时间")
    counts = forecasts.groupby(forecasts["issue_time"].dt.date).size()
    if len(counts) != 365 or not (counts == 4).all():
        raise ValueError("附件3不是每天4次预测")

    audit = {
        "附件1": {
            "工作表": ["Sheet1"],
            "数据形状": list(ref.shape),
            "字段": list(ref.columns),
            "时段": [input_time_labels[0], input_time_labels[-1]],
        },
        "附件2": {
            "工作表": sheets,
            "负荷形状": list(load_raw.shape),
            "光伏形状": list(pv_raw.shape),
            "日期范围": [str(dates.min().date()), str(dates.max().date())],
        },
        "附件3": {
            "工作表": ["Sheet1"],
            "数据形状": list(forecast_raw.shape),
            "日期空白数_填充前": int(forecast_raw.iloc[:, 0].isna().sum()),
            "日期范围": [
                str(forecasts["issue_time"].min()),
                str(forecasts["issue_time"].max()),
            ],
            "每天预测次数": 4,
        },
        "单位处理": "附件1/2/3的功率除以6后进入优化，单位kWh/10分钟",
    }

    return ProblemData(
        dates=dates,
        input_time_labels=input_time_labels,
        price=price,
        reference_load_kwh=reference_load_kwh,
        reference_pv_kwh=reference_pv_kwh,
        actual_load_kwh=actual_load_kwh,
        actual_pv_kwh=actual_pv_kwh,
        forecasts=forecasts,
        data_audit=audit,
    )


def make_targets(decision_time: pd.Timestamp) -> pd.DatetimeIndex:
    """生成决策后每10分钟的目标时刻，最多24小时，年末截断。"""

    end = min(decision_time + pd.Timedelta(hours=24), YEAR_END)
    count = int((end - decision_time).total_seconds() // (SLOT_MINUTES * 60))
    return pd.date_range(
        decision_time + pd.Timedelta(minutes=SLOT_MINUTES),
        periods=count,
        freq=f"{SLOT_MINUTES}min",
    )


def interpolate_pv_forecast(
    data: ProblemData,
    decision_time: pd.Timestamp,
    targets: pd.DatetimeIndex,
) -> np.ndarray:
    """以决策时刻已观测光伏为锚点，将24个整点预测线性插成10分钟预测。"""

    issue = pd.Timestamp(decision_time)
    if issue not in data._forecast_lookup:
        raise ValueError(f"附件3缺少预测：{issue}")
    anchor = data.actual_at(issue, "pv", missing=0.0)
    if not np.isfinite(anchor):
        anchor = 0.0
    hourly_kw = np.asarray(data._forecast_lookup[issue], dtype=float)
    anchor_minutes = np.arange(0, 25, dtype=float) * 60.0
    anchor_values = np.concatenate(([anchor * 6.0], hourly_kw))
    target_minutes = np.asarray(
        [(pd.Timestamp(t) - issue).total_seconds() / 60.0 for t in targets],
        dtype=float,
    )
    interpolated_kw = np.interp(target_minutes, anchor_minutes, anchor_values)
    return np.maximum(interpolated_kw, 0.0) / 6.0


def forecast_load(
    data: ProblemData,
    decision_time: pd.Timestamp,
    targets: pd.DatetimeIndex,
    history_days: int,
) -> np.ndarray:
    """仅用决策日以前的数据形成负荷曲线，并用当天已发生数据修正短期偏差。"""

    day = decision_time.date()
    prior_indices = [idx for idx, value in enumerate(data.dates) if value.date() < day]
    if prior_indices:
        chosen = prior_indices[-history_days:]
        profile = data.actual_load_kwh[chosen].mean(axis=0)
    else:
        profile = data.reference_load_kwh.copy()
    result = np.asarray([profile[data.slot_for_target(t)] for t in targets], dtype=float)

    minutes_now = decision_time.hour * 60 + decision_time.minute
    observed_count = minutes_now // SLOT_MINUTES
    if observed_count > 0 and day in data._date_to_index:
        day_idx = data.day_index(day)
        start = max(0, observed_count - 18)
        residual = data.actual_load_kwh[day_idx, start:observed_count] - profile[start:observed_count]
        if residual.size:
            bias = float(np.median(residual))
            lead = np.arange(1, len(result) + 1, dtype=float)
            result = result + bias * np.exp(-lead / 36.0)
    return np.maximum(result, 0.0)


def build_scenarios(
    data: ProblemData,
    decision_time: pd.Timestamp,
    config: RunConfig,
) -> ScenarioBundle:
    """Walk-forward抽取同一发布时间的成熟24小时联合残差轨迹。"""

    targets = make_targets(decision_time)
    base_load = forecast_load(data, decision_time, targets, config.history_days)
    base_pv = interpolate_pv_forecast(data, decision_time, targets)
    cutoff = decision_time - pd.Timedelta(days=config.history_days)
    candidates = data.forecasts[
        (data.forecasts["issue_time"] < decision_time)
        & (data.forecasts["issue_time"] >= cutoff)
        & (data.forecasts["issue_time"].dt.hour == decision_time.hour)
        & (data.forecasts["issue_time"] + pd.Timedelta(hours=24) <= decision_time)
    ].sort_values("issue_time")

    load_residuals: list[np.ndarray] = []
    pv_residuals: list[np.ndarray] = []
    source_issues: list[pd.Timestamp] = []
    for row in candidates.itertuples(index=False):
        source_issue = pd.Timestamp(row.issue_time)
        source_targets = pd.date_range(
            source_issue + pd.Timedelta(minutes=SLOT_MINUTES),
            periods=SLOTS_PER_DAY,
            freq=f"{SLOT_MINUTES}min",
        )
        actual_load = data.actual_vector(source_targets, "load")
        actual_pv = data.actual_vector(source_targets, "pv")
        if not np.isfinite(actual_load).all() or not np.isfinite(actual_pv).all():
            continue
        historical_load_forecast = forecast_load(
            data, source_issue, source_targets, config.history_days
        )
        historical_pv_forecast = interpolate_pv_forecast(
            data, source_issue, source_targets
        )
        load_residuals.append(actual_load - historical_load_forecast)
        pv_residuals.append(actual_pv - historical_pv_forecast)
        source_issues.append(source_issue)

    if not source_issues:
        scenario_load = base_load[None, :]
        scenario_pv = base_pv[None, :]
        probabilities = np.ones(1)
        selected_sources: list[pd.Timestamp] = []
    else:
        source_count = len(source_issues)
        seed_offset = int((decision_time - YEAR_START).total_seconds() // 3600)
        rng = np.random.default_rng(config.seed + seed_offset)
        if source_count > config.scenario_count:
            selected = np.sort(
                rng.choice(source_count, size=config.scenario_count, replace=False)
            )
        else:
            selected = np.arange(source_count)
        horizon = len(targets)
        scenario_load = np.asarray(
            [np.maximum(base_load + load_residuals[i][:horizon], 0.0) for i in selected]
        )
        scenario_pv = np.asarray(
            [np.maximum(base_pv + pv_residuals[i][:horizon], 0.0) for i in selected]
        )
        selected_sources = [source_issues[i] for i in selected]
        probabilities = np.full(len(selected), 1.0 / len(selected))

    return ScenarioBundle(
        targets=targets,
        base_load=base_load,
        base_pv=base_pv,
        load=scenario_load,
        pv=scenario_pv,
        probabilities=probabilities,
        source_issues=selected_sources,
        decision_time=decision_time,
    )


class SparseRows:
    """按坐标增量构造稀疏约束矩阵。"""

    def __init__(self) -> None:
        self.rows: list[int] = []
        self.cols: list[int] = []
        self.values: list[float] = []
        self.rhs: list[float] = []

    def add(self, coefficients: Iterable[tuple[int, float]], rhs: float) -> None:
        row = len(self.rhs)
        for column, value in coefficients:
            if value:
                self.rows.append(row)
                self.cols.append(int(column))
                self.values.append(float(value))
        self.rhs.append(float(rhs))

    def matrix(self, columns: int):
        return coo_matrix(
            (self.values, (self.rows, self.cols)),
            shape=(len(self.rhs), columns),
        ).tocsr()


def normalize_strategy(name: str) -> str:
    value = STRATEGY_ALIASES.get(name, name)
    if value not in STRATEGIES:
        raise ValueError(f"未知策略{name}，可选值：{sorted(STRATEGIES)}")
    return value


def next_update_slot(strategy: str, decision_slot: int) -> int:
    future = [UPDATE_SLOTS[h] for h in STRATEGIES[strategy] if UPDATE_SLOTS[h] > decision_slot]
    return min(future) if future else SLOTS_PER_DAY


def _terminal_value(price: np.ndarray, params: StorageParams) -> float:
    """一单位SOC未来可替代eta_d单位购电，取预测窗电价中位数估值。"""

    return float(params.eta_discharge * np.median(price))


def solve_stochastic_mpc(
    bundle: ScenarioBundle,
    price: np.ndarray,
    params: StorageParams,
    config: RunConfig,
    strategy: str,
    decision_slot: int,
    soc_start: float,
    g0_fixed: np.ndarray | None,
) -> PlanningResult:
    """求解两阶段补救近似。

    0:00时，全天g0为场景共享变量；允许未来更新覆盖的时段使用
    场景相关调整作为补救。日内节点只把下一更新节点前的购电量设为
    场景共享变量，更远期方案是场景相关补救，不会直接提交。
    """

    strategy = normalize_strategy(strategy)
    horizon = len(bundle.targets)
    scenario_count = bundle.net_load.shape[0]
    current_day_count = min(SLOTS_PER_DAY - decision_slot, horizon)
    commit_end = next_update_slot(strategy, decision_slot)
    commit_count = min(commit_end - decision_slot, current_day_count)
    initial_decision = decision_slot == 0

    cursor = 0

    def allocate(count: int) -> np.ndarray:
        nonlocal cursor
        result = np.arange(cursor, cursor + count, dtype=int)
        cursor += count
        return result

    g0_index = allocate(current_day_count) if initial_decision else np.array([], dtype=int)
    commit_index = allocate(commit_count) if not initial_decision else np.array([], dtype=int)
    scenario_indices: list[dict[str, np.ndarray]] = []
    for _ in range(scenario_count):
        scenario_indices.append(
            {
                "g": allocate(horizon),
                "charge": allocate(horizon),
                "discharge": allocate(horizon),
                "soc": allocate(horizon),
                "emergency": allocate(horizon),
                "unused": allocate(horizon),
                "down": allocate(current_day_count),
                "up": allocate(current_day_count),
            }
        )
    variable_count = cursor
    objective = np.zeros(variable_count)
    bounds: list[tuple[float | None, float | None]] = [(0.0, None)] * variable_count

    if initial_decision:
        objective[g0_index] = price[:current_day_count]
    for indices in scenario_indices:
        bounds_for_charge = [(0.0, params.max_charge_kwh)] * horizon
        bounds_for_discharge = [(0.0, params.max_discharge_kwh)] * horizon
        for idx, bound in zip(indices["charge"], bounds_for_charge):
            bounds[idx] = bound
        for idx, bound in zip(indices["discharge"], bounds_for_discharge):
            bounds[idx] = bound
        for idx in indices["soc"]:
            bounds[idx] = (params.soc_min, params.soc_max)

    equality = SparseRows()
    inequality = SparseRows()
    active_update_slots = [UPDATE_SLOTS[h] for h in STRATEGIES[strategy]]
    final_hard_terminal = bool(bundle.targets[-1] == YEAR_END)
    salvage = _terminal_value(price, params)

    for scenario, (indices, probability) in enumerate(
        zip(scenario_indices, bundle.probabilities)
    ):
        objective[indices["emergency"]] = (
            probability * config.emergency_multiplier * price
        )
        objective[indices["charge"]] = probability * config.throughput_penalty
        objective[indices["discharge"]] = probability * config.throughput_penalty
        if final_hard_terminal:
            equality.add([(indices["soc"][-1], 1.0)], config.final_soc)
        else:
            objective[indices["soc"][-1]] -= probability * salvage

        for h in range(horizon):
            equality.add(
                [
                    (indices["g"][h], 1.0),
                    (indices["discharge"][h], 1.0),
                    (indices["emergency"][h], 1.0),
                    (indices["charge"][h], -1.0),
                    (indices["unused"][h], -1.0),
                ],
                bundle.net_load[scenario, h],
            )
            soc_coefficients = [
                (indices["soc"][h], 1.0),
                (indices["charge"][h], -params.eta_charge),
                (indices["discharge"][h], 1.0 / params.eta_discharge),
            ]
            if h == 0:
                soc_rhs = soc_start
            else:
                soc_coefficients.append((indices["soc"][h - 1], -1.0))
                soc_rhs = 0.0
            equality.add(soc_coefficients, soc_rhs)

            if h >= current_day_count:
                objective[indices["g"][h]] += probability * price[h]

        for h in range(current_day_count):
            target_slot = decision_slot + h
            down_idx = indices["down"][h]
            up_idx = indices["up"][h]
            if initial_decision:
                base_coefficients = [(g0_index[h], -1.0)]
                base_rhs = 0.0
                inequality.add([(down_idx, 1.0), (g0_index[h], -1.0)], 0.0)
                can_adjust = any(slot <= target_slot for slot in active_update_slots)
            else:
                if g0_fixed is None:
                    raise ValueError("日内调整必须提供0:00初始计划g0")
                base_coefficients = []
                base_rhs = float(g0_fixed[target_slot])
                bounds[down_idx] = (0.0, base_rhs)
                can_adjust = True

            equality.add(
                [
                    (indices["g"][h], 1.0),
                    (down_idx, 1.0),
                    (up_idx, -1.0),
                    *base_coefficients,
                ],
                base_rhs,
            )
            if not can_adjust:
                bounds[down_idx] = (0.0, 0.0)
                bounds[up_idx] = (0.0, 0.0)
            else:
                objective[down_idx] += probability * (-0.5 * price[h])
                objective[up_idx] += probability * (1.5 * price[h])

            if not initial_decision and h < commit_count:
                equality.add(
                    [(indices["g"][h], 1.0), (commit_index[h], -1.0)], 0.0
                )

    result = linprog(
        c=objective,
        A_ub=inequality.matrix(variable_count) if inequality.rhs else None,
        b_ub=np.asarray(inequality.rhs) if inequality.rhs else None,
        A_eq=equality.matrix(variable_count),
        b_eq=np.asarray(equality.rhs),
        bounds=bounds,
        method="highs",
        options={"dual_feasibility_tolerance": config.solver_tolerance},
    )
    if not result.success:
        raise RuntimeError(
            f"关键随机MPC求解失败：status={result.status}, message={result.message}"
        )

    solution = result.x
    scenario_g = np.vstack([solution[x["g"]] for x in scenario_indices])
    scenario_soc = np.vstack([solution[x["soc"]] for x in scenario_indices])
    scenario_charge = np.vstack([solution[x["charge"]] for x in scenario_indices])
    scenario_discharge = np.vstack([solution[x["discharge"]] for x in scenario_indices])
    proposed = np.average(
        scenario_g[:, :current_day_count], axis=0, weights=bundle.probabilities
    )
    if initial_decision:
        g0 = solution[g0_index]
        committed = g0[:commit_count]
        proposed = g0.copy()
    else:
        g0 = None
        committed = solution[commit_index]
        proposed[:commit_count] = committed

    return PlanningResult(
        objective=float(result.fun),
        g0=g0,
        committed_block=committed,
        proposed_current_day=proposed,
        mean_soc=np.average(scenario_soc, axis=0, weights=bundle.probabilities),
        scenario_g=scenario_g,
        scenario_charge=scenario_charge,
        scenario_discharge=scenario_discharge,
        solver_message=str(result.message),
    )


def make_soc_grid(params: StorageParams, step: float, extra: Iterable[float] = ()) -> np.ndarray:
    if step <= 0:
        raise ValueError("--soc-step必须为正数")
    grid = np.arange(params.soc_min, params.soc_max + 0.5 * step, step)
    grid = np.concatenate((grid, np.asarray([params.soc_initial, *extra], dtype=float)))
    grid = np.unique(np.clip(grid, params.soc_min, params.soc_max))
    return np.sort(grid)


def feedback_value_function(
    committed_g: np.ndarray,
    scenario_net: np.ndarray,
    probabilities: np.ndarray,
    price: np.ndarray,
    grid: np.ndarray,
    params: StorageParams,
    config: RunConfig,
    terminal_target: float,
    hard_terminal: bool,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """构造一个交付块内的因果DP价值函数，场景保持整段时间相关性。"""

    delta = grid[None, :] - grid[:, None]
    feasible = (
        (delta <= params.eta_charge * params.max_charge_kwh + 1.0e-9)
        & (delta >= -params.max_discharge_kwh / params.eta_discharge - 1.0e-9)
    )
    grid_effect = np.where(
        delta >= 0.0, delta / params.eta_charge, params.eta_discharge * delta
    )
    periods = len(committed_g)
    value = np.empty((periods + 1, len(grid)), dtype=float)
    if hard_terminal:
        value[-1] = np.inf
        value[-1, int(np.argmin(np.abs(grid - terminal_target)))] = 0.0
    else:
        value[-1] = config.feedback_terminal_penalty * np.abs(grid - terminal_target)

    throughput = np.where(
        delta >= 0.0,
        delta / params.eta_charge,
        -params.eta_discharge * delta,
    )
    for t in range(periods - 1, -1, -1):
        shortage = np.maximum(
            scenario_net[:, t, None, None]
            + grid_effect[None, :, :]
            - committed_g[t],
            0.0,
        )
        expected_shortage = np.tensordot(probabilities, shortage, axes=(0, 0))
        cost = (
            config.emergency_multiplier * price[t] * expected_shortage
            + config.throughput_penalty * throughput
            + value[t + 1][None, :]
        )
        cost = np.where(feasible, cost, np.inf)
        value[t] = np.min(cost, axis=1)
    return value, grid_effect, feasible


def execute_feedback_block(
    committed_g: np.ndarray,
    actual_net: np.ndarray,
    scenario_net: np.ndarray,
    probabilities: np.ndarray,
    price: np.ndarray,
    soc_start: float,
    terminal_target: float,
    hard_terminal: bool,
    params: StorageParams,
    config: RunConfig,
) -> dict[str, np.ndarray]:
    """当前10分钟真实净负荷到达后选择动作；不采用情景规划动作。"""

    grid = make_soc_grid(params, config.soc_step, [soc_start, terminal_target])
    value, grid_effect, feasible = feedback_value_function(
        committed_g,
        scenario_net,
        probabilities,
        price,
        grid,
        params,
        config,
        terminal_target,
        hard_terminal,
    )
    current = int(np.argmin(np.abs(grid - soc_start)))
    if abs(grid[current] - soc_start) > 1.0e-7:
        raise RuntimeError("当前SOC无法进入反馈DP网格")
    periods = len(committed_g)
    result = {
        name: np.zeros(periods, dtype=float)
        for name in ("soc_start", "soc_end", "charge", "discharge", "emergency", "unused")
    }
    for t in range(periods):
        result["soc_start"][t] = grid[current]
        shortage = np.maximum(
            actual_net[t] + grid_effect[current] - committed_g[t], 0.0
        )
        transition_cost = (
            config.emergency_multiplier * price[t] * shortage
            + config.throughput_penalty
            * np.where(
                grid_effect[current] >= 0.0,
                grid_effect[current],
                -grid_effect[current],
            )
            + value[t + 1]
        )
        transition_cost = np.where(feasible[current], transition_cost, np.inf)
        nxt = int(np.argmin(transition_cost))
        if not np.isfinite(transition_cost[nxt]):
            raise RuntimeError(f"反馈DP在块内第{t}个时段无可行动作")
        delta_soc = grid[nxt] - grid[current]
        if delta_soc >= 0.0:
            charge = delta_soc / params.eta_charge
            discharge = 0.0
        else:
            charge = 0.0
            discharge = -params.eta_discharge * delta_soc
        result["charge"][t] = charge
        result["discharge"][t] = discharge
        result["emergency"][t] = max(
            actual_net[t] + charge - discharge - committed_g[t], 0.0
        )
        result["unused"][t] = max(
            committed_g[t] + discharge - actual_net[t] - charge, 0.0
        )
        current = nxt
        result["soc_end"][t] = grid[current]
    return result


def interval_label(slot: int) -> str:
    """按区间结束时刻口径输出第slot个10分钟交付区间。"""

    start = slot * SLOT_MINUTES
    end = (slot + 1) * SLOT_MINUTES

    def label(minutes: int) -> str:
        if minutes == 24 * 60:
            return "00:00+1"
        return f"{minutes // 60:02d}:{minutes % 60:02d}"

    return f"({label(start)}, {label(end)}]"


def _price_for_targets(data: ProblemData, targets: pd.DatetimeIndex) -> np.ndarray:
    return np.asarray([data.price[data.slot_for_target(t)] for t in targets], dtype=float)


def _decision_slots(strategy: str) -> list[int]:
    return [0, *[UPDATE_SLOTS[h] for h in STRATEGIES[normalize_strategy(strategy)]]]


def _timestamp_for_slot(day: date, slot: int) -> pd.Timestamp:
    return pd.Timestamp(day) + pd.Timedelta(minutes=slot * SLOT_MINUTES)


def simulate_strategy(
    data: ProblemData,
    start_day: date,
    end_day: date,
    strategy: str,
    initial_soc: float,
    params: StorageParams,
    config: RunConfig,
) -> dict[str, object]:
    """逐日、逐决策块执行一个策略；SOC在日期边界连续传递。"""

    strategy = normalize_strategy(strategy)
    if start_day > end_day:
        raise ValueError("start_day不能晚于end_day")
    if not (params.soc_min <= initial_soc <= params.soc_max):
        raise ValueError("initial_soc超出储能上下界")

    detail_rows: list[dict] = []
    ledger_rows: list[dict] = []
    forecast_audit_rows: list[dict] = []
    daily_rows: list[dict] = []
    current_soc = float(initial_soc)
    started = time.perf_counter()

    for day_timestamp in pd.date_range(start_day, end_day, freq="D"):
        day = day_timestamp.date()
        day_idx = data.day_index(day)
        day_soc_start = current_soc
        g0: np.ndarray | None = None
        effective_g: np.ndarray | None = None
        day_detail_start = len(detail_rows)
        decision_slots = _decision_slots(strategy)

        for decision_number, decision_slot in enumerate(decision_slots):
            decision_time = _timestamp_for_slot(day, decision_slot)
            bundle = build_scenarios(data, decision_time, config)
            prices = _price_for_targets(data, bundle.targets)
            planning = solve_stochastic_mpc(
                bundle=bundle,
                price=prices,
                params=params,
                config=config,
                strategy=strategy,
                decision_slot=decision_slot,
                soc_start=current_soc,
                g0_fixed=g0,
            )

            if decision_slot == 0:
                if planning.g0 is None or len(planning.g0) != SLOTS_PER_DAY:
                    raise RuntimeError("0:00求解未返回完整的全天g0")
                g0 = planning.g0.copy()
                effective_g = g0.copy()
            elif g0 is None or effective_g is None:
                raise RuntimeError("日内调整发生在g0建立之前")

            block_end = next_update_slot(strategy, decision_slot)
            block_end = min(block_end, SLOTS_PER_DAY)
            block_count = block_end - decision_slot
            before_commit = effective_g.copy()
            effective_g[decision_slot:block_end] = planning.committed_block[:block_count]

            current_day_count = SLOTS_PER_DAY - decision_slot
            proposed = planning.proposed_current_day[:current_day_count]
            for relative_slot in range(current_day_count):
                target_slot = decision_slot + relative_slot
                target_time = _timestamp_for_slot(day, target_slot + 1)
                is_locked = target_slot < block_end
                ledger_rows.append(
                    {
                        "strategy": strategy,
                        "decision_time": decision_time,
                        "target_time": target_time,
                        "forecast_issue_time": decision_time,
                        "decision_slot": decision_slot,
                        "target_slot": target_slot,
                        "interval": interval_label(target_slot),
                        "g0_kwh": float(g0[target_slot]),
                        "proposed_g_kwh": float(proposed[relative_slot]),
                        "committed_g_kwh": float(effective_g[target_slot]),
                        "previous_committed_g_kwh": float(before_commit[target_slot]),
                        "is_locked_this_decision": bool(is_locked),
                        "scenario_count": int(bundle.net_load.shape[0]),
                        "solver_objective": float(planning.objective),
                    }
                )

            latest_mature_target = (
                max(bundle.source_issues) + pd.Timedelta(hours=24)
                if bundle.source_issues
                else pd.NaT
            )
            forecast_audit_rows.append(
                {
                    "strategy": strategy,
                    "decision_time": decision_time,
                    "forecast_issue_time": decision_time,
                    "scenario_count": int(bundle.net_load.shape[0]),
                    "source_issue_times": "|".join(str(x) for x in bundle.source_issues),
                    "latest_source_target_time": latest_mature_target,
                    "mature_only": bool(
                        pd.isna(latest_mature_target) or latest_mature_target <= decision_time
                    ),
                    "planned_simultaneous_max_kwh": float(
                        np.minimum(
                            planning.scenario_charge, planning.scenario_discharge
                        ).max()
                    ),
                    "solver_status": planning.solver_message,
                }
            )

            actual_net = (
                data.actual_load_kwh[day_idx, decision_slot:block_end]
                - data.actual_pv_kwh[day_idx, decision_slot:block_end]
            )
            scenario_net = bundle.net_load[:, :block_count]
            block_prices = prices[:block_count]
            hard_terminal = day == date(2025, 12, 31) and block_end == SLOTS_PER_DAY
            terminal_target = (
                config.final_soc
                if hard_terminal
                else float(planning.mean_soc[block_count - 1])
            )
            feedback = execute_feedback_block(
                committed_g=effective_g[decision_slot:block_end],
                scenario_net=scenario_net,
                probabilities=bundle.probabilities,
                price=block_prices,
                params=params,
                config=config,
                terminal_target=terminal_target,
                hard_terminal=hard_terminal,
                actual_net=actual_net,
                soc_start=current_soc,
            )

            for offset in range(block_count):
                slot = decision_slot + offset
                p = float(data.price[slot])
                planned = float(g0[slot])
                final_g = float(effective_g[slot])
                downward = max(planned - final_g, 0.0)
                upward = max(final_g - planned, 0.0)
                charge = float(feedback["charge"][offset])
                discharge = float(feedback["discharge"][offset])
                emergency = float(feedback["emergency"][offset])
                unused = float(feedback["unused"][offset])
                base_cost = p * planned
                downward_refund = 0.5 * p * downward
                upward_premium = 1.5 * p * upward
                emergency_cost = config.emergency_multiplier * p * emergency
                throughput_cost = config.throughput_penalty * (charge + discharge)
                settlement_cost = (
                    base_cost - downward_refund + upward_premium + emergency_cost
                )
                total_cost = settlement_cost + throughput_cost
                target_time = _timestamp_for_slot(day, slot + 1)
                load = float(data.actual_load_kwh[day_idx, slot])
                pv = float(data.actual_pv_kwh[day_idx, slot])
                detail_rows.append(
                    {
                        "strategy": strategy,
                        "date": pd.Timestamp(day),
                        "slot": slot,
                        "interval": interval_label(slot),
                        "target_time": target_time,
                        "committing_decision_time": decision_time,
                        "g0_kwh": planned,
                        "final_purchase_kwh": final_g,
                        "downward_adjustment_kwh": downward,
                        "upward_adjustment_kwh": upward,
                        "actual_load_kwh": load,
                        "actual_pv_kwh": pv,
                        "charge_kwh": charge,
                        "discharge_kwh": discharge,
                        "emergency_purchase_kwh": emergency,
                        "curtailment_or_unused_kwh": unused,
                        "soc_start_kwh": float(feedback["soc_start"][offset]),
                        "soc_end_kwh": float(feedback["soc_end"][offset]),
                        "price_yuan_per_kwh": p,
                        "base_purchase_cost_yuan": base_cost,
                        "downward_refund_yuan": downward_refund,
                        "upward_premium_yuan": upward_premium,
                        "emergency_cost_yuan": emergency_cost,
                        "throughput_penalty_yuan": throughput_cost,
                        "settlement_cost_yuan": settlement_cost,
                        "total_cost_yuan": total_cost,
                        "balance_residual_kwh": (
                            final_g + discharge + emergency
                            - charge - unused - load + pv
                        ),
                        "soc_residual_kwh": (
                            float(feedback["soc_end"][offset])
                            - float(feedback["soc_start"][offset])
                            - params.eta_charge * charge
                            + discharge / params.eta_discharge
                        ),
                    }
                )
            current_soc = float(feedback["soc_end"][-1])

        day_frame = pd.DataFrame(detail_rows[day_detail_start:])
        daily_rows.append(
            {
                "strategy": strategy,
                "date": pd.Timestamp(day),
                "soc_start_kwh": day_soc_start,
                "soc_end_kwh": current_soc,
                "base_purchase_cost_yuan": day_frame["base_purchase_cost_yuan"].sum(),
                "downward_refund_yuan": day_frame["downward_refund_yuan"].sum(),
                "upward_premium_yuan": day_frame["upward_premium_yuan"].sum(),
                "adjustment_cost_yuan": (
                    -day_frame["downward_refund_yuan"].sum()
                    + day_frame["upward_premium_yuan"].sum()
                ),
                "emergency_cost_yuan": day_frame["emergency_cost_yuan"].sum(),
                "throughput_penalty_yuan": day_frame["throughput_penalty_yuan"].sum(),
                "total_cost_yuan": day_frame["total_cost_yuan"].sum(),
                "planned_purchase_kwh": day_frame["g0_kwh"].sum(),
                "final_purchase_kwh": day_frame["final_purchase_kwh"].sum(),
                "emergency_purchase_kwh": day_frame["emergency_purchase_kwh"].sum(),
                "curtailment_or_unused_kwh": day_frame["curtailment_or_unused_kwh"].sum(),
                "storage_throughput_kwh": (
                    day_frame["charge_kwh"].sum() + day_frame["discharge_kwh"].sum()
                ),
            }
        )

    return {
        "strategy": strategy,
        "detail": pd.DataFrame(detail_rows),
        "daily": pd.DataFrame(daily_rows),
        "ledger": pd.DataFrame(ledger_rows),
        "forecast_audit": pd.DataFrame(forecast_audit_rows),
        "runtime_seconds": time.perf_counter() - started,
    }


def validate_run(
    run: dict[str, object],
    params: StorageParams,
    config: RunConfig,
) -> dict[str, object]:
    """执行守恒、成本、时序、边界和信息成熟性检查；失败即抛错。"""

    detail = run["detail"]
    daily = run["daily"]
    ledger = run["ledger"]
    forecast_audit = run["forecast_audit"]
    assert isinstance(detail, pd.DataFrame)
    assert isinstance(daily, pd.DataFrame)
    assert isinstance(ledger, pd.DataFrame)
    assert isinstance(forecast_audit, pd.DataFrame)

    counts = detail.groupby("date").size()
    duplicate_count = int(detail.duplicated(["date", "slot"]).sum())
    numeric = detail.select_dtypes(include=[np.number])
    finite = bool(np.isfinite(numeric.to_numpy()).all())
    nonnegative_columns = [
        "g0_kwh",
        "final_purchase_kwh",
        "downward_adjustment_kwh",
        "upward_adjustment_kwh",
        "charge_kwh",
        "discharge_kwh",
        "emergency_purchase_kwh",
        "curtailment_or_unused_kwh",
        "downward_refund_yuan",
    ]
    min_nonnegative = float(detail[nonnegative_columns].min().min())
    max_balance = float(detail["balance_residual_kwh"].abs().max())
    max_soc_residual = float(detail["soc_residual_kwh"].abs().max())
    max_charge = float(detail["charge_kwh"].max())
    max_discharge = float(detail["discharge_kwh"].max())
    simultaneous = np.minimum(detail["charge_kwh"], detail["discharge_kwh"])
    simultaneous_count = int((simultaneous > 1.0e-7).sum())
    planned_simultaneous_max = float(
        forecast_audit["planned_simultaneous_max_kwh"].max()
    )

    recomputed = (
        detail["base_purchase_cost_yuan"]
        - detail["downward_refund_yuan"]
        + detail["upward_premium_yuan"]
        + detail["emergency_cost_yuan"]
        + detail["throughput_penalty_yuan"]
    )
    max_cost_residual = float((recomputed - detail["total_cost_yuan"]).abs().max())
    soc_bounds_ok = bool(
        (detail["soc_end_kwh"] >= params.soc_min - 1.0e-7).all()
        and (detail["soc_end_kwh"] <= params.soc_max + 1.0e-7).all()
    )
    target_after_decision = bool(
        (pd.to_datetime(ledger["target_time"]) > pd.to_datetime(ledger["decision_time"])).all()
    )
    uncommitted = ledger[~ledger["is_locked_this_decision"].astype(bool)]
    uncommitted_unchanged = bool(
        np.allclose(
            uncommitted["committed_g_kwh"],
            uncommitted["previous_committed_g_kwh"],
            atol=1.0e-8,
            rtol=0.0,
        )
    )
    strategy = str(run["strategy"])
    active_slots = _decision_slots(strategy)
    expected_commit_slot = np.asarray(
        [max(x for x in active_slots if x <= int(slot)) for slot in detail["slot"]],
        dtype=int,
    )
    actual_commit_slot = (
        (
            pd.to_datetime(detail["committing_decision_time"])
            - pd.to_datetime(detail["date"])
        ).dt.total_seconds().to_numpy()
        / (SLOT_MINUTES * 60)
    ).astype(int)
    historical_lock_ok = bool(np.array_equal(expected_commit_slot, actual_commit_slot))
    mature_only = bool(forecast_audit["mature_only"].all())
    continuity_residual = 0.0
    if len(daily) > 1:
        continuity_residual = float(
            np.max(
                np.abs(
                    daily["soc_start_kwh"].to_numpy()[1:]
                    - daily["soc_end_kwh"].to_numpy()[:-1]
                )
            )
        )
    final_required = pd.Timestamp(daily["date"].iloc[-1]).date() == date(2025, 12, 31)
    final_soc_residual = (
        abs(float(daily["soc_end_kwh"].iloc[-1]) - config.final_soc)
        if final_required
        else 0.0
    )

    checks = {
        "strategy": run["strategy"],
        "days": int(len(daily)),
        "detail_rows": int(len(detail)),
        "each_day_144_slots": bool((counts == SLOTS_PER_DAY).all()),
        "duplicate_date_slot_count": duplicate_count,
        "all_numeric_finite": finite,
        "minimum_nonnegative_quantity": min_nonnegative,
        "max_balance_residual_kwh": max_balance,
        "max_soc_recursion_residual_kwh": max_soc_residual,
        "soc_bounds_ok": soc_bounds_ok,
        "max_charge_kwh": max_charge,
        "max_discharge_kwh": max_discharge,
        "simultaneous_actual_count": simultaneous_count,
        "planned_simultaneous_max_kwh": planned_simultaneous_max,
        "max_cost_identity_residual_yuan": max_cost_residual,
        "all_targets_after_decision": target_after_decision,
        "uncommitted_proposals_do_not_change_commitment": uncommitted_unchanged,
        "historical_decision_lock_ok": historical_lock_ok,
        "mature_history_only": mature_only,
        "cross_day_soc_continuity_residual_kwh": continuity_residual,
        "year_end_terminal_required": final_required,
        "year_end_terminal_residual_kwh": final_soc_residual,
        "runtime_seconds": float(run["runtime_seconds"]),
    }
    failures = []
    conditions = {
        "每天144时段": checks["each_day_144_slots"],
        "每时段只结算一次": duplicate_count == 0,
        "数值有限": finite,
        "非负变量": min_nonnegative >= -1.0e-7,
        "功率平衡": max_balance <= 1.0e-6,
        "SOC递推": max_soc_residual <= 1.0e-6,
        "SOC边界": soc_bounds_ok,
        "充电功率边界": max_charge <= params.max_charge_kwh + 1.0e-6,
        "放电功率边界": max_discharge <= params.max_discharge_kwh + 1.0e-6,
        "实际无同时充放": simultaneous_count == 0,
        "成本恒等式": max_cost_residual <= 1.0e-7,
        "交付时段晚于决策": target_after_decision,
        "远期提案未提前锁定": uncommitted_unchanged,
        "历史决策锁定": historical_lock_ok,
        "历史误差均成熟": mature_only,
        "跨日SOC连续": continuity_residual <= 1.0e-7,
        "全年末SOC": final_soc_residual <= config.soc_step / 2 + 1.0e-7,
    }
    for label, passed in conditions.items():
        if not passed:
            failures.append(label)
    checks["all_checks_passed"] = not failures
    checks["failed_checks"] = failures
    if failures:
        raise RuntimeError(f"运行结果检查失败：{failures}")
    return checks


def summarize_run(run: dict[str, object]) -> dict[str, float | str]:
    detail = run["detail"]
    assert isinstance(detail, pd.DataFrame)
    return {
        "strategy": str(run["strategy"]),
        "total_cost_yuan": float(detail["total_cost_yuan"].sum()),
        "base_purchase_cost_yuan": float(detail["base_purchase_cost_yuan"].sum()),
        "adjustment_cost_yuan": float(
            -detail["downward_refund_yuan"].sum()
            + detail["upward_premium_yuan"].sum()
        ),
        "emergency_cost_yuan": float(detail["emergency_cost_yuan"].sum()),
        "emergency_purchase_kwh": float(detail["emergency_purchase_kwh"].sum()),
        "upward_adjustment_kwh": float(detail["upward_adjustment_kwh"].sum()),
        "downward_adjustment_kwh": float(detail["downward_adjustment_kwh"].sum()),
        "curtailment_or_unused_kwh": float(
            detail["curtailment_or_unused_kwh"].sum()
        ),
        "storage_throughput_kwh": float(
            detail["charge_kwh"].sum() + detail["discharge_kwh"].sum()
        ),
        "runtime_seconds": float(run["runtime_seconds"]),
    }


def compute_shapley(comparison: pd.DataFrame) -> pd.DataFrame:
    """由8种子集成本计算6/12/18点信息的Shapley成本节约。"""

    subset_by_strategy = {name: frozenset(hours) for name, hours in STRATEGIES.items()}
    cost = {
        subset_by_strategy[row.strategy]: float(row.total_cost_yuan)
        for row in comparison.itertuples(index=False)
    }
    players = (6, 12, 18)
    if set(cost) != set(subset_by_strategy.values()):
        raise ValueError("Shapley计算需要全部8种更新组合")
    rows = []
    n = len(players)
    for player in players:
        contribution = 0.0
        for subset, base_cost in cost.items():
            if player in subset:
                continue
            k = len(subset)
            weight = math.factorial(k) * math.factorial(n - k - 1) / math.factorial(n)
            saving = base_cost - cost[frozenset((*subset, player))]
            contribution += weight * saving
        rows.append({"forecast_hour": player, "shapley_cost_saving_yuan": contribution})
    return pd.DataFrame(rows)


def _copy_row_style(sheet, source_row: int, target_row: int, max_column: int) -> None:
    for column in range(1, max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy.copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy.copy(source.alignment)
        if source.font:
            target.font = copy.copy(source.font)
        if source.fill:
            target.fill = copy.copy(source.fill)
        if source.border:
            target.border = copy.copy(source.border)


def write_result3(
    template_path: Path,
    output_path: Path,
    detail: pd.DataFrame,
) -> dict[str, object]:
    """保留附件5模板拓扑，填入2月1日至12月31日正式提交结果。"""

    if not template_path.exists():
        raise FileNotFoundError(f"结果模板不存在：{template_path}")
    formal = detail[
        (pd.to_datetime(detail["date"]) >= pd.Timestamp("2025-02-01"))
        & (pd.to_datetime(detail["date"]) <= pd.Timestamp("2025-12-31"))
    ].copy()
    expected_days = pd.date_range("2025-02-01", "2025-12-31", freq="D")
    if not pd.DatetimeIndex(formal["date"].drop_duplicates()).equals(expected_days):
        raise ValueError("生成result3.xlsx需要完整的2月1日至12月31日回放结果")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)
    workbook = load_workbook(output_path)
    required = ["计划购电量", "调整购电量", "充放电量", "紧急购电量"]
    if workbook.sheetnames != required:
        raise ValueError(f"result3模板工作表不符合预期：{workbook.sheetnames}")

    indexed = formal.set_index([pd.to_datetime(formal["date"]).dt.date, "slot"])
    for sheet_name, value_column in (
        ("计划购电量", "g0_kwh"),
        ("调整购电量", "final_purchase_kwh"),
    ):
        sheet = workbook[sheet_name]
        if sheet.max_row != 335 or sheet.max_column != 147:
            raise ValueError(f"{sheet_name}模板尺寸改变：{sheet.max_row}×{sheet.max_column}")
        for row, day in enumerate(expected_days, start=2):
            sheet.cell(row, 1, day.to_pydatetime())
            values = np.asarray(
                [indexed.loc[(day.date(), slot), value_column] for slot in range(144)],
                dtype=float,
            )
            for slot, value in enumerate(values, start=2):
                sheet.cell(row, slot, float(value))
            sheet.cell(row, 146, float(values.sum()))
            if sheet_name == "计划购电量":
                cost = float(
                    formal[formal["date"] == day]["base_purchase_cost_yuan"].sum()
                )
            else:
                rows = formal[formal["date"] == day]
                cost = float(
                    rows["base_purchase_cost_yuan"].sum()
                    - rows["downward_refund_yuan"].sum()
                    + rows["upward_premium_yuan"].sum()
                )
            sheet.cell(row, 147, cost)

    storage = workbook["充放电量"]
    if storage.max_row > 2:
        storage.delete_rows(3, storage.max_row - 2)
    row = 2
    for day in expected_days:
        rows = formal[formal["date"] == day].sort_values("slot")
        for block in range(6):
            if row > 2:
                _copy_row_style(storage, 2, row, 6)
            block_rows = rows.iloc[block * 24 : (block + 1) * 24]
            storage.cell(row, 1, day.to_pydatetime() if block == 0 else None)
            storage.cell(row, 2, f"{block * 4:02d}:00-{(block + 1) * 4:02d}:00")
            storage.cell(row, 3, float(block_rows["charge_kwh"].sum()))
            storage.cell(row, 4, float(block_rows["discharge_kwh"].sum()))
            storage.cell(row, 5, f"{(block + 1) * 4:02d}:00")
            storage.cell(row, 6, float(block_rows["soc_end_kwh"].iloc[-1]))
            row += 1

    emergency_sheet = workbook["紧急购电量"]
    if emergency_sheet.max_row > 2:
        emergency_sheet.delete_rows(3, emergency_sheet.max_row - 2)
    emergency_rows = formal[formal["emergency_purchase_kwh"] > 1.0e-9]
    if emergency_rows.empty:
        emergency_sheet.cell(2, 1, expected_days[0].to_pydatetime())
        emergency_sheet.cell(2, 2, "无")
        emergency_sheet.cell(2, 3, 0.0)
    else:
        for output_row, value in enumerate(emergency_rows.itertuples(index=False), start=2):
            if output_row > 2:
                _copy_row_style(emergency_sheet, 2, output_row, 3)
            emergency_sheet.cell(output_row, 1, pd.Timestamp(value.date).to_pydatetime())
            emergency_sheet.cell(output_row, 2, value.interval)
            emergency_sheet.cell(output_row, 3, float(value.emergency_purchase_kwh))

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(output_path)

    check = load_workbook(output_path, read_only=False, data_only=False)
    required_dates = {
        date(2025, 3, 20), date(2025, 6, 20), date(2025, 9, 20), date(2025, 12, 20)
    }
    date_rows = {
        check["计划购电量"].cell(row, 1).value.date(): row
        for row in range(2, 336)
    }
    missing_plan_cells = sum(
        check["计划购电量"].cell(row, column).value is None
        for row in range(2, 336)
        for column in range(1, 148)
    )
    missing_adjust_cells = sum(
        check["调整购电量"].cell(row, column).value is None
        for row in range(2, 336)
        for column in range(1, 148)
    )
    result = {
        "path": str(output_path),
        "sheetnames": check.sheetnames,
        "plan_shape": [check["计划购电量"].max_row, check["计划购电量"].max_column],
        "adjust_shape": [check["调整购电量"].max_row, check["调整购电量"].max_column],
        "storage_rows": check["充放电量"].max_row,
        "emergency_rows": check["紧急购电量"].max_row,
        "missing_plan_cells": missing_plan_cells,
        "missing_adjust_cells": missing_adjust_cells,
        "required_date_rows": {
            str(value): date_rows.get(value) for value in sorted(required_dates)
        },
        "required_dates_present": required_dates.issubset(date_rows),
    }
    check.close()
    if missing_plan_cells or missing_adjust_cells or not result["required_dates_present"]:
        raise RuntimeError(f"result3.xlsx完整性检查失败：{result}")
    return result


def save_run(run: dict[str, object], output_dir: Path, checks: dict[str, object]) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    for key, filename in (
        ("detail", "problem3_detail.csv"),
        ("daily", "daily_summary.csv"),
        ("ledger", "decision_ledger.csv"),
        ("forecast_audit", "forecast_audit.csv"),
    ):
        frame = run[key]
        assert isinstance(frame, pd.DataFrame)
        frame.to_csv(output_dir / filename, index=False, encoding="utf-8-sig")
    (output_dir / "checks.json").write_text(
        json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def _parse_date(text: str) -> date:
    return pd.Timestamp(text).date()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="问题3三时间尺度因果滚动随机MPC")
    parser.add_argument("--data-dir", type=Path, default=Path("附件"))
    parser.add_argument(
        "--output-dir", "--out-dir", dest="output_dir", type=Path,
        default=Path("analysis_outputs/problem3")
    )
    parser.add_argument("--template", type=Path, default=None)
    parser.add_argument(
        "--strategy", default="S061218",
        help="S0/S06/S12/S18/S0612/S0618/S1218/S061218；兼容S6和Sall"
    )
    parser.add_argument("--date", type=_parse_date, default=date(2025, 1, 1))
    parser.add_argument("--all-year", action="store_true")
    parser.add_argument("--initial-soc", type=float, default=None)
    parser.add_argument("--ablation", action="store_true", help="运行S0/S06/S0612/S061218")
    parser.add_argument("--all-subsets", action="store_true", help="运行全部8种更新组合并计算Shapley")
    parser.add_argument("--scenarios", type=int, default=RunConfig.scenario_count)
    parser.add_argument("--history-days", type=int, default=RunConfig.history_days)
    parser.add_argument("--seed", type=int, default=RunConfig.seed)
    parser.add_argument("--soc-step", type=float, default=RunConfig.soc_step)
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    if args.scenarios <= 0 or args.history_days <= 0:
        raise ValueError("场景数和历史天数必须为正数")
    data = read_problem_data(args.data_dir)
    params = StorageParams()
    config = RunConfig(
        scenario_count=args.scenarios,
        history_days=args.history_days,
        seed=args.seed,
        soc_step=args.soc_step,
    )
    if args.all_year:
        start_day, end_day = date(2025, 1, 1), date(2025, 12, 31)
        initial_soc = params.soc_initial if args.initial_soc is None else args.initial_soc
        run_mode = "formal_full_year"
    elif args.date == date(2025, 1, 1):
        start_day = end_day = args.date
        initial_soc = params.soc_initial if args.initial_soc is None else args.initial_soc
        run_mode = "jan1_initialization_and_real_operation"
    elif args.initial_soc is None:
        start_day, end_day = date(2025, 1, 1), args.date
        initial_soc = params.soc_initial
        run_mode = "causal_replay_from_jan1"
    else:
        start_day = end_day = args.date
        initial_soc = args.initial_soc
        run_mode = "isolated_diagnostic_with_explicit_initial_soc"

    if not date(2025, 1, 1) <= end_day <= date(2025, 12, 31):
        raise ValueError("日期必须位于2025年")
    if args.all_subsets:
        strategies = list(STRATEGIES)
    elif args.ablation:
        strategies = list(CUMULATIVE_ABLATION)
    else:
        strategies = [normalize_strategy(args.strategy)]

    args.output_dir.mkdir(parents=True, exist_ok=True)
    metadata = {
        "run_mode": run_mode,
        "start_day": str(start_day),
        "end_day": str(end_day),
        "initial_soc_kwh": initial_soc,
        "strategies": strategies,
        "storage": asdict(params),
        "config": asdict(config),
        "data_audit": data.data_audit,
        "template_time_header_discrepancy": (
            "模板首列写0:10-0:20，但数据首槽按强制口径解释为(00:00,00:10]；"
            "程序保留模板标题并按144列顺序映射。"
        ),
    }
    (args.output_dir / "run_config.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    summaries = []
    primary_run: dict[str, object] | None = None
    for strategy in strategies:
        print(f"运行{strategy}: {start_day}至{end_day}", flush=True)
        run = simulate_strategy(
            data=data,
            start_day=start_day,
            end_day=end_day,
            strategy=strategy,
            initial_soc=initial_soc,
            params=params,
            config=config,
        )
        checks = validate_run(run, params, config)
        strategy_dir = args.output_dir if len(strategies) == 1 else args.output_dir / strategy
        save_run(run, strategy_dir, checks)
        summaries.append(summarize_run(run))
        if strategy == normalize_strategy(args.strategy):
            primary_run = run
    comparison = pd.DataFrame(summaries)
    comparison.to_csv(args.output_dir / "strategy_comparison.csv", index=False, encoding="utf-8-sig")
    if args.all_subsets:
        shapley = compute_shapley(comparison)
        shapley.to_csv(args.output_dir / "shapley_values.csv", index=False, encoding="utf-8-sig")

    if args.all_year:
        if primary_run is None:
            primary_run = simulate_strategy(
                data, start_day, end_day, normalize_strategy(args.strategy), initial_soc, params, config
            )
            validate_run(primary_run, params, config)
        template = args.template or args.data_dir / "附件5" / "result3.xlsx"
        workbook_check = write_result3(
            template, args.output_dir / "result3.xlsx", primary_run["detail"]
        )
        (args.output_dir / "workbook_check.json").write_text(
            json.dumps(workbook_check, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    print(comparison.to_string(index=False), flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
